import os
import re
import ast
import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# All 6 scenarios definition based on training_logs
LOG_FILES = [
    {
        "filepath": "training_logs/train_cifar100.txt",
        "scenario_key": "cifar100",
        "scenario_name": "CIFAR-100 (10 tasks)",
        "sheet_name": "CIFAR-100",
        "num_tasks": 10
    },
    {
        "filepath": "training_logs/train_cub200.txt",
        "scenario_key": "cub200",
        "scenario_name": "CUB-200 (10 tasks)",
        "sheet_name": "CUB-200",
        "num_tasks": 10
    },
    {
        "filepath": "training_logs/train_inr100.txt",
        "scenario_key": "imagenetr_10t",
        "scenario_name": "ImageNet-R (10 tasks)",
        "sheet_name": "ImageNet-R (10T)",
        "num_tasks": 10
    },
    {
        "filepath": "training_logs/train_ina.txt",
        "scenario_key": "imageneta_20t",
        "scenario_name": "ImageNet-A (20 tasks)",
        "sheet_name": "ImageNet-A",
        "num_tasks": 20
    },
    {
        "filepath": "training_logs/train_long_1.txt",
        "scenario_key": "imagenetr_20t",
        "scenario_name": "ImageNet-R (20 tasks)",
        "sheet_name": "ImageNet-R (20T)",
        "num_tasks": 20
    },
    {
        "filepath": "training_logs/train_long_2.txt",
        "scenario_key": "imagenetr_40t",
        "scenario_name": "ImageNet-R (40 tasks)",
        "sheet_name": "ImageNet-R (40T)",
        "num_tasks": 40
    }
]

METHOD_MAP = [
    ("exp_finetune", "Aper-Finetune"),
    ("exp_vpt-shallow", "Aper-VPT-Shallow"),
    ("exp_vpt-deep", "Aper-VPT-Deep"),
    ("exp_adapter", "Aper-Adapter"),
    ("exp_ssf", "Aper-SSF")
]

def calc_mean_std(values):
    n = len(values)
    if n == 0:
        return 0.0, 0.0
    mean = sum(values) / n
    if n > 1:
        var = sum((x - mean) ** 2 for x in values) / (n - 1)
        std = math.sqrt(var)
    else:
        std = 0.0
    return mean, std

def parse_log_file(filepath):
    with open(filepath, "r", encoding="utf-8") as f:
        text = f.read()

    sections = re.split(r"\[\d+/5\]\s+Running:\s+Model=", text)
    data_by_method = {}

    for sec in sections[1:]:
        header_line = sec.split("\n")[0]
        prefix_match = re.search(r"\((\S+)\)", header_line)
        if not prefix_match:
            continue
        prefix = prefix_match.group(1)
        method_name = dict(METHOD_MAP).get(prefix, prefix)

        seed_blocks = re.split(r"\[trainer\.py\]\s+=>\s+seed:\s*(\d+)", sec)
        seeds_data = []

        for i in range(1, len(seed_blocks), 2):
            s_num = int(seed_blocks[i])
            s_body = seed_blocks[i+1]

            cnn_dicts = [ast.literal_eval(m) for m in re.findall(r"CNN:\s*(\{[^\n]+\})", s_body)]
            if not cnn_dicts:
                continue

            T = len(cnn_dicts)
            acc_curve = [d["total"] for d in cnn_dicts]
            news = [d["new"] for d in cnn_dicts]

            # Standard CL Forgetting: drop of previously learned tasks compared to their initial accuracy
            fgt_curve = [0.0]
            for k in range(1, T):
                avg_initial = sum(news[:k]) / k
                drop = avg_initial - cnn_dicts[k]["old"]
                fgt_curve.append(drop)

            avg_inc_acc = sum(acc_curve) / len(acc_curve)
            final_acc = acc_curve[-1]

            seeds_data.append({
                "seed": s_num,
                "T": T,
                "acc_curve": acc_curve,
                "fgt_curve": fgt_curve,
                "avg_inc_acc": avg_inc_acc,
                "final_acc": final_acc
            })

        if seeds_data:
            T = seeds_data[0]["T"]
            mean_acc_tasks = [sum(s["acc_curve"][t] for s in seeds_data) / len(seeds_data) for t in range(T)]
            mean_fgt_tasks = [sum(s["fgt_curve"][t] for s in seeds_data) / len(seeds_data) for t in range(T)]

            avg_incs = [s["avg_inc_acc"] for s in seeds_data]
            final_accs = [s["final_acc"] for s in seeds_data]

            m_avg_inc, s_avg_inc = calc_mean_std(avg_incs)
            m_final, s_final = calc_mean_std(final_accs)

            data_by_method[method_name] = {
                "seeds": seeds_data,
                "mean_acc_tasks": mean_acc_tasks,
                "mean_fgt_tasks": mean_fgt_tasks,
                "avg_inc_mean": m_avg_inc,
                "avg_inc_std": s_avg_inc,
                "final_mean": m_final,
                "final_std": s_final
            }

    return data_by_method

def build_excel(output_filename="results.xlsx"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # Typography & Styles
    font_title = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="595959")
    font_section = Font(name="Calibri", size=12, bold=True, color="1F4E79")
    font_header_main = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_header_sub = Font(name="Calibri", size=10, bold=True, color="1F4E79")
    font_data = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)

    fill_header_main = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_header_sub = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_alt_row = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='C0C0C0'),
        right=Side(style='thin', color='C0C0C0'),
        top=Side(style='thin', color='C0C0C0'),
        bottom=Side(style='thin', color='C0C0C0')
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. Parse all scenarios
    scenarios_data = {}
    for sc in LOG_FILES:
        scenarios_data[sc["scenario_key"]] = {
            "config": sc,
            "data": parse_log_file(sc["filepath"])
        }

    # =========================================================================
    # SHEET 1: "Tổng hợp"
    # =========================================================================
    ws_summary = wb.create_sheet(title="Tổng hợp")
    ws_summary.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_summary["A1"] = "BẢNG TỔNG HỢP KẾT QUẢ THỰC NGHIỆM 5 BIẾN THỂ APER"
    ws_summary["A1"].font = font_title
    ws_summary["A2"] = "Các chỉ số Average Incremental Accuracy và Final Accuracy được tính trung bình (mean) và độ lệch chuẩn (std) qua 3 seed [0, 1, 2]"
    ws_summary["A2"].font = font_subtitle

    # Table 1: Matrix comparison across all 6 scenarios
    start_row = 4
    ws_summary.row_dimensions[start_row].height = 26
    ws_summary.row_dimensions[start_row+1].height = 22

    # Column A: Phương pháp
    ws_summary.merge_cells(start_row=start_row, start_column=1, end_row=start_row+1, end_column=1)
    cell_method = ws_summary.cell(row=start_row, column=1, value="Phương pháp")
    cell_method.font = font_header_main
    cell_method.fill = fill_header_main
    cell_method.alignment = align_center
    cell_method.border = thin_border
    ws_summary.cell(row=start_row+1, column=1).border = thin_border
    ws_summary.cell(row=start_row+1, column=1).fill = fill_header_main

    col_idx = 2
    for sc in LOG_FILES:
        sc_name = sc["scenario_name"]
        ws_summary.merge_cells(start_row=start_row, start_column=col_idx, end_row=start_row, end_column=col_idx+1)
        top_cell = ws_summary.cell(row=start_row, column=col_idx, value=sc_name)
        top_cell.font = font_header_main
        top_cell.fill = fill_header_main
        top_cell.alignment = align_center
        top_cell.border = thin_border
        ws_summary.cell(row=start_row, column=col_idx+1).border = thin_border
        ws_summary.cell(row=start_row, column=col_idx+1).fill = fill_header_main

        # Subheaders
        c1 = ws_summary.cell(row=start_row+1, column=col_idx, value="Average Incremental Acc (%)")
        c2 = ws_summary.cell(row=start_row+1, column=col_idx+1, value="Final Acc (%)")
        for c in (c1, c2):
            c.font = font_header_sub
            c.fill = fill_header_sub
            c.alignment = align_center
            c.border = thin_border

        col_idx += 2

    # Fill data rows
    row_curr = start_row + 2
    for idx_m, (_, method_name) in enumerate(METHOD_MAP):
        ws_summary.row_dimensions[row_curr].height = 20
        c_m = ws_summary.cell(row=row_curr, column=1, value=method_name)
        c_m.font = font_bold
        c_m.alignment = align_left
        c_m.border = thin_border
        if idx_m % 2 == 1:
            c_m.fill = fill_alt_row

        col_idx = 2
        for sc in LOG_FILES:
            sc_key = sc["scenario_key"]
            m_data = scenarios_data[sc_key]["data"].get(method_name)
            if m_data:
                inc_str = f"{m_data['avg_inc_mean']:.2f} ± {m_data['avg_inc_std']:.2f}"
                fin_str = f"{m_data['final_mean']:.2f} ± {m_data['final_std']:.2f}"
            else:
                inc_str = "N/A"
                fin_str = "N/A"

            c_inc = ws_summary.cell(row=row_curr, column=col_idx, value=inc_str)
            c_fin = ws_summary.cell(row=row_curr, column=col_idx+1, value=fin_str)
            for c in (c_inc, c_fin):
                c.font = font_data
                c.alignment = align_center
                c.border = thin_border
                if idx_m % 2 == 1:
                    c.fill = fill_alt_row

            col_idx += 2
        row_curr += 1

    # Detailed Table: Breakdown by Scenario and Individual Seeds
    row_curr += 2
    ws_summary.cell(row=row_curr, column=1, value="BẢNG CHI TIẾT SỐ LIỆU TỪNG SEED THEO KỊCH BẢN").font = font_section
    row_curr += 2

    for sc in LOG_FILES:
        sc_name = sc["scenario_name"]
        sc_key = sc["scenario_key"]

        ws_summary.cell(row=row_curr, column=1, value=f"• Kịch bản: {sc_name}").font = font_header_sub
        row_curr += 1

        headers_detail = [
            "Phương pháp",
            "Avg Inc Acc\n(Seed 0)", "Avg Inc Acc\n(Seed 1)", "Avg Inc Acc\n(Seed 2)", "Mean Avg Inc Acc\n(%)", "Std Avg Inc Acc\n(%)",
            "Final Acc\n(Seed 0)", "Final Acc\n(Seed 1)", "Final Acc\n(Seed 2)", "Mean Final Acc\n(%)", "Std Final Acc\n(%)"
        ]
        ws_summary.row_dimensions[row_curr].height = 28
        for c_idx, h_text in enumerate(headers_detail, 1):
            cell = ws_summary.cell(row=row_curr, column=c_idx, value=h_text)
            cell.font = font_header_sub
            cell.fill = fill_header_sub
            cell.alignment = align_center
            cell.border = thin_border

        row_curr += 1
        for idx_m, (_, method_name) in enumerate(METHOD_MAP):
            ws_summary.row_dimensions[row_curr].height = 20
            m_data = scenarios_data[sc_key]["data"].get(method_name)
            c_m = ws_summary.cell(row=row_curr, column=1, value=method_name)
            c_m.font = font_bold
            c_m.alignment = align_left
            c_m.border = thin_border
            if idx_m % 2 == 1:
                c_m.fill = fill_alt_row

            if m_data and len(m_data["seeds"]) == 3:
                s0 = m_data["seeds"][0]
                s1 = m_data["seeds"][1]
                s2 = m_data["seeds"][2]
                vals = [
                    round(s0["avg_inc_acc"], 2),
                    round(s1["avg_inc_acc"], 2),
                    round(s2["avg_inc_acc"], 2),
                    round(m_data["avg_inc_mean"], 2),
                    round(m_data["avg_inc_std"], 2),
                    round(s0["final_acc"], 2),
                    round(s1["final_acc"], 2),
                    round(s2["final_acc"], 2),
                    round(m_data["final_mean"], 2),
                    round(m_data["final_std"], 2)
                ]
            else:
                vals = ["N/A"] * 10

            for c_idx, v in enumerate(vals, 2):
                cell = ws_summary.cell(row=row_curr, column=c_idx, value=v)
                cell.font = font_data
                cell.alignment = align_right if isinstance(v, (int, float)) else align_center
                cell.border = thin_border
                if isinstance(v, (int, float)):
                    cell.number_format = '0.00'
                if idx_m % 2 == 1:
                    cell.fill = fill_alt_row

            row_curr += 1
        row_curr += 2

    # Column dimensions for Summary Sheet
    ws_summary.column_dimensions['A'].width = 22
    for c_idx in range(2, 2 + len(LOG_FILES) * 2):
        col_letter = get_column_letter(c_idx)
        ws_summary.column_dimensions[col_letter].width = 18

    # Freeze panes on summary sheet
    ws_summary.freeze_panes = "B6"

    # =========================================================================
    # SHEETS 2 - 7: TỪNG KỊCH BẢN (ACCURACY & FORGETTING TỪNG TASK)
    # =========================================================================
    for sc in LOG_FILES:
        sheet_title = sc["sheet_name"]
        sc_key = sc["scenario_key"]
        sc_name = sc["scenario_name"]
        num_tasks = sc["num_tasks"]
        data_methods = scenarios_data[sc_key]["data"]

        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        # Sheet header title
        ws["A1"] = f"KỊCH BẢN: {sc_name.upper()} - KẾT QUẢ TỪNG TASK"
        ws["A1"].font = font_title
        ws["A2"] = "Các chỉ số Accuracy (%) và Forgetting (%) của từng task được tính trung bình (mean) qua 3 seed [0, 1, 2]"
        ws["A2"].font = font_subtitle

        start_r = 4
        ws.row_dimensions[start_r].height = 26
        ws.row_dimensions[start_r+1].height = 20

        # Method column header (Row 4-5 merged)
        ws.merge_cells(start_row=start_r, start_column=1, end_row=start_r+1, end_column=1)
        cell_m = ws.cell(row=start_r, column=1, value="Phương pháp")
        cell_m.font = font_header_main
        cell_m.fill = fill_header_main
        cell_m.alignment = align_center
        cell_m.border = thin_border
        ws.cell(row=start_r+1, column=1).border = thin_border
        ws.cell(row=start_r+1, column=1).fill = fill_header_main

        # Task headers
        for t in range(num_tasks):
            t_col_start = 2 + t * 2
            t_col_end = t_col_start + 1
            ws.merge_cells(start_row=start_r, start_column=t_col_start, end_row=start_r, end_column=t_col_end)
            t_cell = ws.cell(row=start_r, column=t_col_start, value=f"Task {t+1}")
            t_cell.font = font_header_main
            t_cell.fill = fill_header_main
            t_cell.alignment = align_center
            t_cell.border = thin_border
            ws.cell(row=start_r, column=t_col_end).border = thin_border
            ws.cell(row=start_r, column=t_col_end).fill = fill_header_main

            # Sub-headers (Accuracy, Forgetting)
            c_acc = ws.cell(row=start_r+1, column=t_col_start, value="Accuracy")
            c_fgt = ws.cell(row=start_r+1, column=t_col_end, value="Forgetting")
            for c in (c_acc, c_fgt):
                c.font = font_header_sub
                c.fill = fill_header_sub
                c.alignment = align_center
                c.border = thin_border

        # Data rows
        curr_row = start_r + 2
        for idx_m, (_, method_name) in enumerate(METHOD_MAP):
            ws.row_dimensions[curr_row].height = 20
            m_cell = ws.cell(row=curr_row, column=1, value=method_name)
            m_cell.font = font_bold
            m_cell.alignment = align_left
            m_cell.border = thin_border
            if idx_m % 2 == 1:
                m_cell.fill = fill_alt_row

            m_data = data_methods.get(method_name)
            for t in range(num_tasks):
                col_acc = 2 + t * 2
                col_fgt = col_acc + 1

                if m_data and t < len(m_data["mean_acc_tasks"]):
                    val_acc = float(round(m_data["mean_acc_tasks"][t], 2))
                    val_fgt = float(round(m_data["mean_fgt_tasks"][t], 2))
                else:
                    val_acc = "N/A"
                    val_fgt = "N/A"

                cell_a = ws.cell(row=curr_row, column=col_acc, value=val_acc)
                cell_f = ws.cell(row=curr_row, column=col_fgt, value=val_fgt)

                for c in (cell_a, cell_f):
                    c.font = font_data
                    c.alignment = align_right if isinstance(c.value, (int, float)) else align_center
                    c.border = thin_border
                    if isinstance(c.value, (int, float)):
                        c.number_format = '0.00'
                    if idx_m % 2 == 1:
                        c.fill = fill_alt_row

            curr_row += 1

        # Freeze panes so Method column and Headers stay visible
        ws.freeze_panes = "B6"

        # Widths
        ws.column_dimensions['A'].width = 20
        for col_idx in range(2, 2 + num_tasks * 2):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 13

    wb.save(output_filename)
    print(f"Generated {output_filename} successfully with sheets: {wb.sheetnames}")

if __name__ == "__main__":
    build_excel("results.xlsx")
